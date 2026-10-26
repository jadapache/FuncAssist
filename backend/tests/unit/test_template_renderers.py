"""Unit tests for declarative DOCX and XLSX template renderers."""

import io
import pytest
from docx import Document
import openpyxl

from app.services.template_renderer import DocxTemplateRenderer, XlsxTemplateRenderer


def test_docx_template_renderer_single_and_split_runs() -> None:
    doc = Document()
    p1 = doc.add_paragraph("Documento: {{artifact.title}}")
    p2 = doc.add_paragraph()
    # Simulate Word run-splitting across multiple runs
    p2.add_run("Problema detectado: {{artifact.")
    p2.add_run("problem}}")

    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].text = "Actor: {{artifact.actor}}"

    buf = io.BytesIO()
    doc.save(buf)
    template_bytes = buf.getvalue()

    data = {
        "artifact": {
            "title": "Sistema de Turnos",
            "problem": "Demora en sala de espera",
            "actor": "Enfermero Jefe",
        }
    }

    rendered_bytes = DocxTemplateRenderer.render(template_bytes, data)
    assert len(rendered_bytes) > 0

    rendered_doc = Document(io.BytesIO(rendered_bytes))
    full_text = " ".join(p.text for p in rendered_doc.paragraphs)
    assert "Documento: Sistema de Turnos" in full_text
    assert "Problema detectado: Demora en sala de espera" in full_text
    assert "Actor: Enfermero Jefe" in rendered_doc.tables[0].rows[0].cells[0].text


def test_xlsx_template_renderer_preserves_formulas() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriz"
    ws["A1"] = "Proyecto: {{artifact.project}}"
    ws["B1"] = "Total Horas:"
    ws["C1"] = "=SUM(10, 20)"

    buf = io.BytesIO()
    wb.save(buf)
    template_bytes = buf.getvalue()

    data = {
        "artifact": {
            "project": "FuncAssist Core",
        }
    }

    rendered_bytes = XlsxTemplateRenderer.render(
        template_bytes, data, cell_coordinates={"D1": "Versión {{artifact.project}}"}
    )
    assert len(rendered_bytes) > 0

    rendered_wb = openpyxl.load_workbook(io.BytesIO(rendered_bytes), data_only=False)
    sheet = rendered_wb["Matriz"]
    assert sheet["A1"].value == "Proyecto: FuncAssist Core"
    assert sheet["C1"].value == "=SUM(10, 20)"
    assert sheet["D1"].value == "Versión FuncAssist Core"
