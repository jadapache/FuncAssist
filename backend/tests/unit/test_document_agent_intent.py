"""Unit tests for document agent dynamic intent classification, multimodal messages, and docx extraction."""

from __future__ import annotations

import io
from docx import Document

from app.api.document_agent import classify_prompt_intent
from app.ai.runner import _assemble_multimodal_message
from app.context.knowledge import _extract_docx, extract_text


def test_classify_prompt_intent_slash_commands():
    assert classify_prompt_intent("/plan Pruebas de integración") == "generic_plan"
    assert classify_prompt_intent("/plan") == "generic_plan"
    assert classify_prompt_intent("/informe Hallazgos de QA") == "generic_report"
    assert classify_prompt_intent("/report Estado de regresión") == "generic_report"


def test_classify_prompt_intent_qa_and_template():
    assert classify_prompt_intent("¿Cómo puedo estructurar una prueba de carga?") == "chat"
    assert classify_prompt_intent("Que cambio hiciste en ese doc?") == "chat"
    assert classify_prompt_intent("¿Qué cambios se hicieron?") == "chat"
    assert classify_prompt_intent("Explica la diferencia entre pruebas unitarias e integración") == "chat"
    assert classify_prompt_intent("Con base al documento adjunto genera informe que esboce resumen de lo requerimientos y funcionalidades detalladas") == "generic_report"
    assert classify_prompt_intent("Genera el documento de mejora para login") == "template_mutation"
    assert classify_prompt_intent("Completa la plantilla corporativa con estos hallazgos") == "template_mutation"


def test_classify_prompt_intent_explicit_mode():
    assert classify_prompt_intent("cualquier texto", mode="plan") == "generic_plan"
    assert classify_prompt_intent("cualquier texto", mode="report") == "generic_report"
    assert classify_prompt_intent("cualquier texto", mode="chat") == "chat"
    assert classify_prompt_intent("cualquier texto", mode="template") == "template_mutation"


def test_assemble_multimodal_message_text_only():
    msg = _assemble_multimodal_message("Consulta simple", None)
    assert msg.role == "user"
    assert msg.content == "Consulta simple"


def test_assemble_multimodal_message_with_images():
    images = [
        {"name": "captura_error.png", "data": "iVBORw0KGgoAAAANSUhEUg==", "mime_type": "image/png"}
    ]
    msg = _assemble_multimodal_message("Analiza esta captura", images)
    assert msg.role == "user"
    assert isinstance(msg.content, list)
    assert len(msg.content) == 2
    assert msg.content[0] == {"type": "text", "text": "Analiza esta captura"}
    assert msg.content[1]["type"] == "image_url"
    assert msg.content[1]["image_url"]["url"].startswith("data:image/png;base64,")


def test_extract_docx_with_paragraphs_and_tables():
    doc = Document()
    doc.add_heading("Especificación Funcional", level=1)
    doc.add_paragraph("Este es un párrafo de requerimientos.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Descripción"
    table.rows[1].cells[0].text = "RF-01"
    table.rows[1].cells[1].text = "Login seguro"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extracted = _extract_docx(docx_bytes)
    assert "Especificación Funcional" in extracted
    assert "Este es un párrafo de requerimientos." in extracted
    assert "RF-01" in extracted
    assert "Login seguro" in extracted
    assert "| ID | Descripción |" in extracted


def test_extract_text_image_safeguard():
    img_bytes = b"\x89PNG\r\n\x1a\nfakeimagecontent"
    result = extract_text("evidencia_error.png", img_bytes)
    assert "[Asset Imagen: evidencia_error.png" in result
