"""Declarative XLSX template renderer with formula and style preservation."""

from __future__ import annotations

import io
from typing import Any, Dict, Optional, Union
import openpyxl


class XlsxTemplateRenderer:
    """Renders structured artifact data into XLSX templates preserving cell styles and formulas."""

    @classmethod
    def render(
        cls,
        template_source: Union[str, bytes, io.BytesIO],
        data: Dict[str, Any],
        cell_coordinates: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        wb = cls._load_workbook(template_source)
        flattened = cls._flatten_dict(data)

        cls._replace_markers_in_sheets(wb, flattened)
        if cell_coordinates:
            cls._apply_explicit_coordinates(wb, cell_coordinates, flattened)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def _load_workbook(source: Union[str, bytes, io.BytesIO]) -> openpyxl.Workbook:
        if isinstance(source, bytes):
            return openpyxl.load_workbook(io.BytesIO(source), data_only=False)
        if isinstance(source, io.BytesIO):
            return openpyxl.load_workbook(source, data_only=False)
        return openpyxl.load_workbook(source, data_only=False)

    @classmethod
    def _replace_markers_in_sheets(cls, wb: openpyxl.Workbook, data: Dict[str, str]) -> None:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    cls._process_cell(cell, data)

    @staticmethod
    def _process_cell(cell: Any, data: Dict[str, str]) -> None:
        if not isinstance(cell.value, str) or "{{" not in cell.value:
            return

        val = cell.value
        for key, replacement in data.items():
            pattern = f"{{{{{key}}}}}"
            if pattern in val:
                val = val.replace(pattern, str(replacement))
        cell.value = val

    @staticmethod
    def _apply_explicit_coordinates(
        wb: openpyxl.Workbook, coords: Dict[str, Any], data: Dict[str, str]
    ) -> None:
        sheet = wb.active
        for coordinate, val in coords.items():
            resolved = str(val)
            for k, v in data.items():
                resolved = resolved.replace(f"{{{{{k}}}}}", str(v))
            sheet[coordinate] = resolved

    @classmethod
    def _flatten_dict(cls, d: Dict[str, Any], prefix: str = "") -> Dict[str, str]:
        items: Dict[str, str] = {}
        for k, v in d.items():
            key = f"{prefix}.{k}" if prefix else k
            if isinstance(v, dict):
                items.update(cls._flatten_dict(v, key))
            elif isinstance(v, list):
                if v and isinstance(v[0], dict):
                    continue
                items[key] = ", ".join(str(item) for item in v)
            else:
                items[key] = str(v) if v is not None else ""
        return items
