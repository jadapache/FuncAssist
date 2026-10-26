"""Unit tests for declarative document agent exports and rubric evaluation."""

from __future__ import annotations

import io
from pathlib import Path
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
import openpyxl

from app.api.document_agent import (
    router,
    _try_parse_structured_artifact,
    _evaluate_artifacts_rubric,
    _resolve_export_data,
)
from app.schemas.artifacts import (
    FunctionalImprovementArtifact,
    FunctionalInventoryArtifact,
    GherkinScenario,
    InventoryModule,
    InventoryComponent,
)


@pytest.fixture
def client() -> TestClient:
    app = FastAPI()
    app.include_router(router, prefix="/api")
    return TestClient(app)


def test_try_parse_structured_artifact_improvement() -> None:
    json_payload = (
        '{\n'
        '  "artifact_id": "art-101",\n'
        '  "title": "Optimización Login",\n'
        '  "problem_statement": "El login actual tiene latencia alta y bloqueos",\n'
        '  "business_objective": "Reducir el tiempo de respuesta a menos de 200ms",\n'
        '  "actors": ["Usuario"],\n'
        '  "main_flow": ["Paso 1", "Paso 2", "Paso 3"],\n'
        '  "alternate_flows": ["Error de red"],\n'
        '  "acceptance_criteria": []\n'
        '}'
    )
    parsed = _try_parse_structured_artifact(json_payload, "Test")
    assert isinstance(parsed, FunctionalImprovementArtifact)
    assert parsed.title == "Optimización Login"


def test_try_parse_structured_artifact_markdown_fences() -> None:
    fenced_payload = (
        '```json\n'
        '{\n'
        '  "artifact_id": "art-202",\n'
        '  "title": "Inventario Core",\n'
        '  "project_name": "FuncAssist",\n'
        '  "modules": []\n'
        '}\n'
        '```'
    )
    parsed = _try_parse_structured_artifact(fenced_payload, "Test")
    assert isinstance(parsed, FunctionalInventoryArtifact)
    assert parsed.project_name == "FuncAssist"


def test_evaluate_artifacts_rubric_approved() -> None:
    art = FunctionalImprovementArtifact(
        artifact_id="art-test",
        title="Validación Biométrica",
        problem_statement="El sistema actual carece de doble factor biométrico en sedes críticas.",
        business_objective="Garantizar la identidad y autenticación estricta del personal.",
        actors=["Médico", "Enfermera"],
        main_flow=["Ingresar cédula", "Capturar huella", "Confirmar acceso"],
        alternate_flows=["Falla en sensor biométrico"],
        acceptance_criteria=[
            GherkinScenario(
                title="Huella válida",
                given=["un médico registrado"],
                when=["coloca su dedo en el lector"],
                then=["se autoriza el ingreso al sistema"],
            )
        ],
    )
    art_dict = {
        "title": art.title,
        "content": art.model_dump_json(),
        "extension": "docx",
    }
    score, feedback = _evaluate_artifacts_rubric([art_dict])
    assert score >= 75.0
    assert "aprobados" in feedback.lower()


def test_export_docx_fallback_markdown(client: TestClient) -> None:
    payload = {
        "title": "Documento_Prueba",
        "markdown": "# Título\n\nEste es el contenido de prueba para exportación.",
    }
    response = client.post("/api/doc-agent/export-docx", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(response.content) > 1000


def test_export_docx_with_template(client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    doc = Document()
    doc.add_paragraph("Hola {{artifact.title}}")
    doc.add_paragraph("Detalle: {{artifact.summary}}")
    tmpl_path = tmp_path / "test_template.docx"
    doc.save(tmpl_path)

    monkeypatch.setattr(
        "app.api.document_agent.get_template_file_path",
        lambda tid: tmpl_path if tid == "tmpl-custom-docx" else None,
    )

    payload = {
        "title": "Sistema Clínico",
        "templateId": "tmpl-custom-docx",
        "artifactData": {
            "artifact": {
                "title": "Sistema Clínico FCV",
                "summary": "Resumen ejecutivo del sistema",
            }
        },
    }
    response = client.post("/api/doc-agent/export-docx", json=payload)
    assert response.status_code == 200
    res_doc = Document(io.BytesIO(response.content))
    all_text = " ".join(p.text for p in res_doc.paragraphs)
    assert "Sistema Clínico FCV" in all_text
    assert "Resumen ejecutivo del sistema" in all_text


def test_export_xlsx_fallback_markdown(client: TestClient) -> None:
    payload = {
        "title": "Matriz_QA",
        "markdown": "| ID | Módulo | Estado |\n|---|---|---|\n| 1 | Auth | OK |",
    }
    response = client.post("/api/doc-agent/export-xlsx", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(response.content) > 1000


def test_export_xlsx_with_template(client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["A1"] = "Proyecto: {{artifact.project}}"
    sheet["A2"] = "Responsable: {{artifact.lead}}"
    tmpl_path = tmp_path / "test_template.xlsx"
    wb.save(tmpl_path)

    monkeypatch.setattr(
        "app.api.document_agent.get_template_file_path",
        lambda tid: tmpl_path if tid == "tmpl-custom-xlsx" else None,
    )

    payload = {
        "title": "Inventario_Prueba",
        "templateId": "tmpl-custom-xlsx",
        "artifactData": {
            "artifact": {
                "project": "Plataforma HIS",
                "lead": "Ingeniero QA",
            }
        },
    }
    response = client.post("/api/doc-agent/export-xlsx", json=payload)
    assert response.status_code == 200
    res_wb = openpyxl.load_workbook(io.BytesIO(response.content))
    res_sheet = res_wb.active
    assert "Plataforma HIS" in res_sheet["A1"].value
    assert "Ingeniero QA" in res_sheet["A2"].value
